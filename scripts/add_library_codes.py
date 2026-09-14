import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CODE_FILE = '/root/.claude/uploads/926f45d0-f134-567b-b05c-3fe0159e34e0/5716d4e8-2026___________________.xlsx'
MAIL_FILE = '/home/user/project_nld/data/대학도서관_설문대상_발송명단.xlsx'

# ---------- load code source: 학교명 -> [(도서관명, 도서관부호), ...] in file order ----------
wb_code = openpyxl.load_workbook(CODE_FILE, data_only=True)
ws_code = wb_code['공시대상학교 리스트']
rows = list(ws_code.iter_rows(min_row=2, values_only=True))

groups = {}
order_per_school = []
cur = None
for r in rows:
    school = r[2]
    if school:
        cur = school
        groups.setdefault(cur, [])
    lib_name, code = r[9], r[10]
    if lib_name:
        groups[cur].append((lib_name.strip(), code))

# ---------- open mailing list and add columns ----------
wb = openpyxl.load_workbook(MAIL_FILE)
ws = wb['대상기관 발송명단']

header_fill = PatternFill(start_color='2B4A73', end_color='2B4A73', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)
thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

N_COL = 14  # column N
ws.cell(row=1, column=N_COL, value='통합도서관부호')
c = ws.cell(row=1, column=N_COL)
c.fill = header_fill
c.font = header_font
c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.column_dimensions['N'].width = 60

filled_single = 0
filled_merged = 0
missing_school = 0

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    school = row[1].value  # column B
    libs = groups.get(school)
    if libs is None:
        missing_school += 1
        continue
    if not libs:
        continue  # no KERIS library row at all for this school

    # representative = first library that has a code; fall back to first entry
    rep_idx = next((i for i, (n, c) in enumerate(libs) if c is not None), 0)
    rep_name, rep_code = libs[rep_idx]

    if rep_code is not None:
        row[3].value = rep_code  # column D 도서관부호

    remaining = [lib for i, lib in enumerate(libs) if i != rep_idx]
    if remaining:
        parts = []
        for name, code in remaining:
            parts.append(f"{name}: {code if code is not None else '코드 미확인'}")
        row[13].value = '; '.join(parts)  # column N
        filled_merged += 1
    else:
        filled_single += 1

    for cell in (row[3], row[13]):
        cell.border = border
        cell.alignment = Alignment(vertical='center', wrap_text=True)

print(f'단일 도서관(대표부호만 입력): {filled_single}')
print(f'통합 도서관(대표+통합도서관부호 입력): {filled_merged}')
print(f'매칭 실패 학교: {missing_school}')

# ---------- update 안내 sheet ----------
note_ws = wb['안내']
note_rows = list(note_ws.iter_rows(min_row=2, max_row=note_ws.max_row))
for r in note_rows:
    if r[0].value == '도서관부호 / 주소':
        r[1].value = (
            '국립중앙도서관 도서관코드 조회 결과를 반영해 도서관부호(D열)를 채웠습니다. '
            '캠퍼스 도서관이 여러 개로 통합된 학교(35개교)는 대표 캠퍼스 1개의 부호만 D열에 넣고, '
            '나머지 캠퍼스 도서관명과 부호는 N열(통합도서관부호)에 "도서관명: 부호" 형식으로 나열했습니다. '
            'KERIS 통계 자체가 없던 학교(16개교) 등 부호를 확인하지 못한 경우는 공란으로 두었습니다. '
            '주소는 이번 작업 범위가 아니어서 공란으로 유지했습니다.'
        )
        break

note_ws.column_dimensions['B'].width = 100

wb.save(MAIL_FILE)
print('saved', MAIL_FILE)
