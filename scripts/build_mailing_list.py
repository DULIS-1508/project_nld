import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

m = json.load(open('/home/user/project_nld/data/master.json'))

MERGE_NOTES = {
    '경남도립거창대학': "통폐합: 구 '경남도립거창대학' → 2026.3. 국립창원대학교로 통합",
    '경남도립남해대학': "통폐합: 구 '경남도립남해대학' → 2026.3. 국립창원대학교로 통합",
    '서라벌대학교': "통폐합: 구 '서라벌대학교' → 2024.3. 경주대학교와 통합, 신경주대학교로 개편",
    '원광보건대학교': "통폐합: 구 '원광보건대학교' → 2026.2. 원광대학교로 흡수통합",
    '국립강릉원주대학교': "통폐합: 구 '국립강릉원주대학교' → 2026.3. 강원대학교와 통합('통합 강원대')",
    '전남도립대학교': "통폐합: 구 '전남도립대학교' → 2026.3. 국립목포대학교와 통합(2·4년제 통합모델)",
}
BRANCH_SCHOOLS = {'건국대학교(글로컬)', '고려대학교(세종)', '동국대학교(WISE)', '연세대학교(미래)', '한양대학교(ERICA)'}

def disability_note(dc_list):
    if not dc_list:
        return '장애학생지원센터 통계: 조사대상 아님(전문대학·대학원대학 등)'
    has_o = any(d.get('지원센터유무') == 'O' for d in dc_list)
    has_x = any(d.get('지원센터유무') == 'X' for d in dc_list)
    if has_o:
        n = sum((d.get('장애재학생수_학부') or 0) for d in dc_list)
        tag = '장애학생지원센터 통계: 있음(O)'
        if len(dc_list) > 1:
            tag += f' (캠퍼스 {len(dc_list)}개 분할집계 중 일부 포함)'
        tag += f' · 장애재학생수(학부) 합계 {n}명'
        return tag
    if has_x:
        return '장애학생지원센터 통계: 없음(X)'
    return '장애학생지원센터 통계: 확인 필요'

rows = []
no_counter = [0]
for name, v in m.items():
    libs = v['libraries']
    merge_notes_found = []
    for lib in libs:
        mm = lib.get('match_method', '')
        if mm.startswith('override:'):
            prefix = mm.split(':', 1)[1]
            note = MERGE_NOTES.get(prefix)
            if note and note not in merge_notes_found:
                merge_notes_found.append(note)
    if not merge_notes_found and v['본분교구분'] == '분교' and name in BRANCH_SCHOOLS:
        merge_notes_found.append('교육부 고시상 본교와 별도의 독립 학교(분교)로 등재')
    if v.get('비고'):
        merge_notes_found.append(f"[교육부고시 비고] {v['비고']}")
    merge_note = ' | '.join(merge_notes_found) if merge_notes_found else None

    dis_note = disability_note(v['disability_center'])

    def make_row(lib):
        no_counter[0] += 1
        remarks = []
        if merge_note:
            remarks.append(merge_note)
        if len(libs) > 1:
            idx = libs.index(lib) + 1
            remarks.append(f'대학도서관 통계 {len(libs)}개 캠퍼스 중 {idx}번째 (구분 집계)')
        remarks.append(dis_note)
        return {
            'no': no_counter[0],
            '대학교명': name,
            '도서관명': lib['lib_row_name'] if lib else '',
            '도서관부호': '',
            '주소': '',
            '본분교구분': lib['본분교'] if lib else v['본분교구분'],
            '대학구분': lib['대학구분'] if lib else v['대학구분'],
            '학교구분': lib['학교구분'] if lib else v['학교구분'],
            '설립구분': lib['설립구분'] if lib else v['설립구분'],
            '근거법령': v['근거법령'],
            '지역': lib['지역'] if lib else v['지역'],
            '시군구': lib.get('시군구', '') if lib else '',
            '비고': ' / '.join(remarks),
        }

    if libs:
        for lib in libs:
            rows.append(make_row(lib))
    else:
        remarks = []
        if merge_note:
            remarks.append(merge_note)
        remarks.append('KERIS 대학도서관 통계 없음')
        remarks.append(dis_note)
        no_counter[0] += 1
        rows.append({
            'no': no_counter[0], '대학교명': name, '도서관명': '', '도서관부호': '', '주소': '',
            '본분교구분': v['본분교구분'], '대학구분': v['대학구분'], '학교구분': v['학교구분'],
            '설립구분': v['설립구분'], '근거법령': v['근거법령'], '지역': v['지역'], '시군구': '',
            '비고': ' / '.join(remarks),
        })

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '대상기관 발송명단'
headers = ['no','대학교명','도서관명','도서관부호','주소','본분교구분','대학구분','학교구분','설립구분','근거법령','지역','시군구','비고']
ws.append(headers)
for r in rows:
    ws.append([r[h] for h in headers])

header_fill = PatternFill(start_color='2B4A73', end_color='2B4A73', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)
thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.freeze_panes = 'A2'
widths = {'A':6,'B':22,'C':28,'D':12,'E':24,'F':11,'G':10,'H':12,'I':10,'J':16,'K':8,'L':10,'M':60}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical='center', wrap_text=(cell.column_letter=='M'))

note_ws = wb.create_sheet('안내')
note_ws.append(['항목', '설명'])
for cell in note_ws[1]:
    cell.fill = header_fill; cell.font = header_font
notes = [
    ('행 구성', '학교당 대학도서관 캠퍼스 통계가 여러 건이면(예: 가톨릭대 3개) 각 캠퍼스 도서관을 별도 행으로 나열했습니다. 도서관 통계가 없는 학교(주로 전문대학·대학원대학)는 도서관명 공란으로 1행 처리했습니다.'),
    ('도서관부호 / 주소', '업로드하신 3개 원본 파일(교육부 고시, KERIS 대학도서관 통계, 장애학생지원센터 현황)에 해당 항목이 없어 공란으로 두었습니다. 별도 자료가 있으면 알려주시면 채워 넣겠습니다.'),
    ('비고 - 통폐합', '최근 통폐합·명칭변경이 확인된 학교는 구 명칭과 현재 명칭, 통합 시기를 표시했습니다(웹 검색으로 사실관계 확인).'),
    ('비고 - 장애학생지원센터', '장애학생지원센터 현황조사 결과를 있음(O)/없음(X)/조사대상 아님으로 표시하고, 있음인 경우 장애재학생수 합계를 함께 표기했습니다.'),
]
for r in notes:
    note_ws.append(r)
note_ws.column_dimensions['A'].width = 20
note_ws.column_dimensions['B'].width = 100
for row in note_ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)

out = '/home/user/project_nld/data/대학도서관_설문대상_발송명단.xlsx'
wb.save(out)
print('saved', out, 'rows=', len(rows))
