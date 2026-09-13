import openpyxl, re, json
from collections import defaultdict

BASE = '/root/.claude/uploads/926f45d0-f134-567b-b05c-3fe0159e34e0'
F1 = f'{BASE}/c6f35fee-2026_____________.xlsx'
F2 = f'{BASE}/dd71e917-________________20260913131948.xlsx'
F3 = f'{BASE}/f2753d93-_______________________2026-09-13132114428.xlsx'

SPECIAL_KEYWORDS = {
    '건국대학교': ('건국대학교(글로컬)', '글로컬'),
    '고려대학교': ('고려대학교(세종)', '세종'),
    '동국대학교': ('동국대학교(WISE)', 'WISE'),
    '연세대학교': ('연세대학교(미래)', '미래'),
    '한양대학교': ('한양대학교(ERICA)', 'ERICA'),
}

# Manually verified overrides (2026 mergers/renames + name-format mismatches
# between KERIS and the MOE gazette). Source: web search, see chat log.
#   - 경남도립거창대학/경남도립남해대학 -> 2026.3 국립창원대학교로 통합
#   - 서라벌대학교 -> 2024.3 경주대학교와 통합, 신경주대학교로 개편
#   - 원광보건대학교 -> 2026.2 원광대학교로 흡수통합
#   - 한국폴리텍 각 캠퍼스 -> 로마숫자 유니코드/띄어쓰기 표기 차이로 인한 매칭 실패
#   - 국립강릉원주대학교 -> 2026.3 강원대학교와 통합('통합 강원대')
#   - 전남도립대학교 -> 2026.3 국립목포대학교와 통합(2·4년제 통합모델)
OVERRIDE_STARTSWITH = [
    ('경남도립거창대학', '국립창원대학교'),
    ('경남도립남해대학', '국립창원대학교'),
    ('서라벌대학교', '신경주대학교'),
    ('원광보건대학교', '원광대학교'),
    ('국립강릉원주대학교', '강원대학교'),
    ('전남도립대학교', '국립목포대학교'),
    ('한국폴리텍Ⅰ대학', '한국폴리텍 I 대학'),
    ('한국폴리텍Ⅱ대학', '한국폴리텍 II 대학'),
    ('한국폴리텍Ⅲ대학', '한국폴리텍 III 대학'),
    ('한국폴리텍Ⅳ대학', '한국폴리텍 IV 대학'),
    ('한국폴리텍Ⅴ대학', '한국폴리텍 V 대학'),
    ('한국폴리텍Ⅵ대학', '한국폴리텍 VI 대학'),
    ('한국폴리텍Ⅶ대학', '한국폴리텍 VII 대학'),
    ('한국폴리텍특성화', '한국폴리텍 특성화대학'),
]
# Rows that are genuinely out of MOE-gazette scope (verified via web search) -
# never force-matched, just recorded for the review log.
KNOWN_OUT_OF_SCOPE = {
    '국제예술대학교': '평생교육법상 학교형태 평생교육시설(전공대학) - 고등교육법 제2조 학교 아님, 공시대상 아님',
    '경남과학기술대학교': '2021년 경상대학교와 통합되어 경상국립대학교로 개편 - file3에 옛 명칭 잔존 데이터',
}

def strip_paren(name):
    return re.sub(r'\(.*?\)', '', name).strip()

# ---------- Load file1 (master) ----------
wb1 = openpyxl.load_workbook(F1, data_only=True, read_only=True)
ws1 = wb1['공시대상학교 리스트']
rows1 = list(ws1.iter_rows(min_row=2, values_only=True))
f1_names = [r[0] for r in rows1]
f1_exact_set = set(f1_names)
base_to_names = defaultdict(list)
for n in f1_names:
    base_to_names[strip_paren(n)].append(n)

master = {}
for r in rows1:
    name = r[0]
    master[name] = {
        'school_name': name,
        '본분교구분': r[1], '대학구분': r[2], '학교구분': r[3],
        '지역': r[4], '설립구분': r[5], '근거법령': r[6], '비고': r[7],
        'libraries': [], 'disability_center': [],
    }

def resolve_school(raw_name_for_exact, raw_full_string_for_keyword):
    """Try exact match, then base match w/ special-case disambiguation."""
    if raw_name_for_exact in f1_exact_set:
        return raw_name_for_exact, 'exact'
    base = strip_paren(raw_name_for_exact)
    candidates = base_to_names.get(base, [])
    if len(candidates) == 1:
        return candidates[0], 'base_unique'
    if len(candidates) > 1:
        # ambiguous - use keyword from the full string
        if base in SPECIAL_KEYWORDS:
            branch_name, keyword = SPECIAL_KEYWORDS[base]
            if keyword in raw_full_string_for_keyword:
                return branch_name, 'keyword_branch'
            else:
                return base, 'keyword_main'  # main campus, exact name == base for these 5
        return None, 'ambiguous_unhandled'
    return None, 'no_match'

# ---------- file2: KERIS library stats ----------
wb2 = openpyxl.load_workbook(F2, data_only=True, read_only=True)
ws2 = wb2['Sheet1']
rows2 = list(ws2.iter_rows(min_row=2, values_only=True))

sorted_full_names = sorted(f1_exact_set, key=len, reverse=True)

def match_lib_row(lib_name):
    # tier 0: manual overrides (mergers/renames verified via web search)
    for prefix, target in OVERRIDE_STARTSWITH:
        if lib_name.startswith(prefix) and target in f1_exact_set:
            return target, f'override:{prefix}'
    # tier 1: literal prefix match against full (parenthesized) names, longest wins
    for fn in sorted_full_names:
        if lib_name.startswith(fn):
            return fn, 'full_prefix'
    # tier 2: base-name prefix match
    # try longest base names as prefixes
    sorted_bases = sorted(base_to_names.keys(), key=len, reverse=True)
    for b in sorted_bases:
        if lib_name.startswith(b):
            resolved, how = resolve_school(b, lib_name)
            if resolved:
                return resolved, f'base_prefix:{how}'
            return None, 'ambiguous_unhandled'
    for known, reason in KNOWN_OUT_OF_SCOPE.items():
        if lib_name.startswith(known):
            return None, f'out_of_scope:{reason}'
    return None, 'no_match'

lib_unmatched = []
lib_matched_count = 0
for r in rows2:
    lib_name = r[0]
    if not lib_name:
        continue
    resolved, how = match_lib_row(lib_name)
    if resolved:
        master[resolved]['libraries'].append({
            'lib_row_name': lib_name, '본분교': r[1], '대학구분': r[2],
            '학교구분': r[3], '지역': r[4], '시군구': r[5], '설립구분': r[6],
            'match_method': how,
        })
        lib_matched_count += 1
    else:
        lib_unmatched.append((lib_name, r[1], how))

print(f'file2 total={len([r for r in rows2 if r[0]])}, matched={lib_matched_count}, unmatched={len(lib_unmatched)}')
print('--- file2 unmatched ---')
for x in lib_unmatched:
    print(' ', x)

# ---------- file3: disability support center census ----------
wb3 = openpyxl.load_workbook(F3, data_only=True, read_only=True)
ws3 = wb3['Sheet1']
rows3 = list(ws3.iter_rows(min_row=4, values_only=True))

d3_unmatched = []
d3_matched_count = 0
i = 0
while i < len(rows3):
    r = rows3[i]
    name_full = r[5]
    if name_full:
        left, _, tag = name_full.partition('_')
        resolved = None
        how = None
        for prefix, target in OVERRIDE_STARTSWITH:
            if left.startswith(prefix) and target in f1_exact_set:
                resolved, how = target, f'override:{prefix}'
                break
        if not resolved:
            resolved, how = resolve_school(left, name_full)
        if not resolved:
            # tier fallback: try full literal prefix against f1 names (rare)
            for fn in sorted_full_names:
                if left.startswith(fn) or fn.startswith(left):
                    resolved, how = fn, 'fuzzy_fallback'
                    break
        if not resolved:
            for known, reason in KNOWN_OUT_OF_SCOPE.items():
                if left.startswith(known):
                    how = f'out_of_scope:{reason}'
                    break
        grad_row = rows3[i+1] if i+1 < len(rows3) else None

        def dis_total(row):
            if row is None:
                return None
            heavy = row[19] or 0   # 계-중증
            mild = row[20] or 0    # 계-경증
            return heavy + mild

        if resolved:
            master[resolved]['disability_center'].append({
                'raw_name': name_full, 'campus_tag': tag,
                '기준연도': r[0], '학교종류': r[1], '설립구분': r[2], '지역': r[3], '상태': r[4],
                '전체재학생수_학부': r[7],
                '장애재학생수_학부': dis_total(r),
                '장애재학생수_대학원': dis_total(grad_row),
                '장애인특별전형유무': r[21] if len(r) > 21 else None,
                '특별지원위원회유무': r[23] if len(r) > 23 else None,
                '지원센터유무': r[24] if len(r) > 24 else None,
                '지원센터_전담인력수': r[27] if len(r) > 27 else None,
                'match_method': how,
            })
            d3_matched_count += 1
        else:
            d3_unmatched.append((name_full, how))
        i += 2
    else:
        i += 1

print(f'\nfile3 total_schools={sum(1 for r in rows3 if r[5])}, matched={d3_matched_count}, unmatched={len(d3_unmatched)}')
print('--- file3 unmatched ---')
for x in d3_unmatched:
    print(' ', x)

json.dump(master, open('/home/user/project_nld/data/master.json','w'), ensure_ascii=False, indent=1, default=str)
print('\nsaved master.json with', len(master), 'schools')
